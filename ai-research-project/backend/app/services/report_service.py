"""Statistical report generation service."""

import io
import logging
from typing import Optional, List, Dict, Any

from app.models.session import AnalysisSession
from app.models.file import UploadedFile as UploadedFileModel

logger = logging.getLogger(__name__)


def generate_word_report_bytes(
    session: AnalysisSession,
    db_file: Optional[UploadedFileModel],
    results: Optional[List[Dict[str, Any]]] = None,
) -> io.BytesIO:
    """Generate a Word (.docx) statistical report.

    Args:
        session: Analysis session model
        db_file: Uploaded file info (optional)
        results: List of analysis result items (optional)

    Returns:
        BytesIO buffer containing the .docx file
    """
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # --- Styles ---
    style = doc.styles["Normal"]
    style.font.name = "SimSun"
    style.font.size = Pt(11)

    # --- Title ---
    title = doc.add_heading("医学统计分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- Meta Info ---
    doc.add_paragraph(f"会话标题: {session.title}")
    doc.add_paragraph(f"创建时间: {session.created_at.strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(f"报告生成时间: {session.updated_at.strftime('%Y-%m-%d %H:%M')}")
    if db_file:
        doc.add_paragraph(
            f"数据文件: {db_file.filename} ({db_file.row_count} 行 × {len(db_file.columns)} 列)"
        )
        doc.add_paragraph(f"数据列: {', '.join(db_file.columns)}")

    doc.add_paragraph("─" * 60)

    # --- Methods ---
    doc.add_heading("统计方法", level=1)
    doc.add_paragraph(
        "本报告使用 DeepSeek AI 辅助统计分析平台自动生成。"
        "分析方法包括描述性统计、组间比较、回归分析等标准医学统计方法。"
        "所有分析均使用 Python (pandas, scipy, statsmodels) 执行。"
    )

    # --- Results ---
    doc.add_heading("分析结果", level=1)

    if results:
        for i, item in enumerate(results, 1):
            item_type = item.get("type", "text")
            title_text = item.get("title", f"结果 {i}")
            content = item.get("content", "")

            if item_type == "text" and content:
                doc.add_heading(title_text, level=2)
                para = doc.add_paragraph(content)
                para.style.font.size = Pt(10)
            elif item_type == "table" and item.get("rows"):
                doc.add_heading(title_text, level=2)
                columns = item.get("columns", [])
                rows = item.get("rows", [])
                if columns and rows:
                    table = doc.add_table(rows=len(rows) + 1, cols=len(columns), style="Light Grid Accent 1")
                    # Header
                    for j, col in enumerate(columns):
                        cell = table.rows[0].cells[j]
                        cell.text = str(col.get("title", "") if isinstance(col, dict) else col)
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                    # Data
                    for i_row, row in enumerate(rows):
                        for j_col, col_def in enumerate(columns):
                            key = col_def.get("key", "") if isinstance(col_def, dict) else col_def
                            val = row.get(key, "")
                            table.rows[i_row + 1].cells[j_col].text = str(val) if val is not None else ""
    else:
        doc.add_paragraph("详细分析结果请参见分析工作台中的输出。")
        doc.add_paragraph("本报告为分析会话的结构化摘要。")

    # --- Footer ---
    doc.add_paragraph("─" * 60)
    footer_para = doc.add_paragraph("报告由 AI 医学统计分析平台自动生成")
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def generate_pdf_report_bytes(
    session: AnalysisSession,
    db_file: Optional[UploadedFileModel],
    results: Optional[List[Dict[str, Any]]] = None,
) -> io.BytesIO:
    """Generate a PDF statistical report.

    Args:
        session: Analysis session model
        db_file: Uploaded file info (optional)
        results: List of analysis result items (optional)

    Returns:
        BytesIO buffer containing the PDF file
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle,
    )

    buf = io.BytesIO()

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=f"医学统计分析报告 - Session {session.id}",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle", parent=styles["Title"],
        fontSize=18, textColor=HexColor("#1a1a1a"),
        spaceAfter=12 * mm,
    )
    heading_style = ParagraphStyle(
        "CustomHeading", parent=styles["Heading1"],
        fontSize=14, textColor=HexColor("#333333"),
        spaceBefore=8 * mm, spaceAfter=4 * mm,
    )
    body_style = ParagraphStyle(
        "CustomBody", parent=styles["Normal"],
        fontSize=10, leading=16,
    )

    story = []

    # Title
    story.append(Paragraph("医学统计分析报告", title_style))
    story.append(HRFlowable(width="100%", thickness=1, color=HexColor("#cccccc")))
    story.append(Spacer(1, 6 * mm))

    # Meta
    story.append(Paragraph(f"<b>会话标题:</b> {session.title}", body_style))
    story.append(Paragraph(
        f"<b>创建时间:</b> {session.created_at.strftime('%Y-%m-%d %H:%M')}", body_style
    ))
    if db_file:
        story.append(Paragraph(
            f"<b>数据文件:</b> {db_file.filename} "
            f"({db_file.row_count} 行 × {len(db_file.columns)} 列)",
            body_style,
        ))
    story.append(Spacer(1, 8 * mm))

    # Methods
    story.append(Paragraph("统计方法", heading_style))
    story.append(Paragraph(
        "本报告使用 DeepSeek AI 辅助统计分析平台自动生成。"
        "分析方法包括描述性统计、组间比较、回归分析等标准医学统计方法。"
        "所有分析均使用 Python (pandas, scipy, statsmodels) 执行。",
        body_style,
    ))
    story.append(Spacer(1, 6 * mm))

    # Results
    story.append(Paragraph("分析结果", heading_style))

    if results:
        for item in results:
            item_type = item.get("type", "text")
            title_text = item.get("title", "")
            content = item.get("content", "")

            if title_text:
                story.append(Paragraph(f"<b>{title_text}</b>", body_style))
            if item_type == "text" and content:
                story.append(Paragraph(content.replace("\n", "<br/>"), body_style))
            elif item_type == "table" and item.get("rows"):
                cols = item.get("columns", [])
                rows = item.get("rows", [])
                if cols and rows:
                    col_keys = [c.get("key", "") if isinstance(c, dict) else c for c in cols]
                    col_titles = [c.get("title", k) if isinstance(c, dict) else k for c, k in zip(cols, col_keys)]
                    table_data = [col_titles]
                    for row in rows:
                        table_data.append([str(row.get(k, "")) for k in col_keys])
                    t = Table(table_data)
                    t.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), HexColor("#e8e8e8")),
                        ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#cccccc")),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ]))
                    story.append(t)
            story.append(Spacer(1, 3 * mm))
    else:
        story.append(Paragraph("详细分析结果请参见分析工作台中的输出。", body_style))

    story.append(Spacer(1, 10 * mm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=HexColor("#cccccc")))
    story.append(Paragraph("<i>报告由 AI 医学统计分析平台自动生成</i>", body_style))

    doc.build(story)
    buf.seek(0)
    return buf
